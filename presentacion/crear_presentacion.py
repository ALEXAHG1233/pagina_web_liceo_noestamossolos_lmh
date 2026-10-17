"""Crea la presentación "No estamos solos LMH" (.pptx) para importarla a Canva.
Paleta Green Screen · Títulos: Josefin Sans · Texto: Libre Franklin."""
import math
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

AQUI = os.path.dirname(os.path.abspath(__file__))
REC = os.path.join(AQUI, "recursos")
SALIDA = os.path.join(AQUI, "No_estamos_solos_LMH.pptx")

BLANCO, MENTA, LIMA, OLIVA, OSCURO = "F4FFFC", "91EAAF", "C3E956", "4D7111", "1F4B2C"
TEXTO_SUAVE = "D3ECDC"
TITULO, CUERPO = "Josefin Sans", "Libre Franklin"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
VACIA = prs.slide_layouts[6]
TOTAL = 27
numero = [0]
fondo_actual = [""]


def rgb(h):
    return RGBColor.from_string(h)


def rec(nombre):
    return os.path.join(REC, nombre)


# ---------------------------------------------------------------------------
# Piezas básicas
# ---------------------------------------------------------------------------
def nueva(fondo="fondo_1.jpg", icono=None, hojas=True):
    s = prs.slides.add_slide(VACIA)
    s.shapes.add_picture(rec(fondo), 0, 0, prs.slide_width, prs.slide_height)
    numero[0] += 1
    if hojas and not fondo.startswith("fondo_liceo"):
        # Hojas 3D asomándose por los bordes, detrás del contenido (como en la referencia).
        izq, der = ("obj_hoja_menta.png", "obj_hoja_lima.png") if numero[0] % 2 else ("obj_hoja_lima.png", "obj_hoja_oscura.png")
        h1 = s.shapes.add_picture(rec(izq), Inches(-1.8), Inches(3.5), Inches(2.5), Inches(2.5))
        h1.rotation = 20 + (numero[0] * 37) % 40
        h2 = s.shapes.add_picture(rec(der), Inches(12.35), Inches(1.9 + (numero[0] % 3) * 0.5), Inches(2.0), Inches(2.0))
        h2.rotation = -15 - (numero[0] * 23) % 50
    if icono:
        claro = fondo == "fondo_brillante.jpg"
        tam = 1.9 if claro else 1.7
        s.shapes.add_picture(rec(f"obj_{icono}.png"), Inches(13.333 - tam - 0.35), Inches(0.2), Inches(tam), Inches(tam))
    fondo_actual[0] = fondo
    return s


def alfa(forma, valor):
    """Transparencia del relleno (valor 0 a 100, 100 = sólido)."""
    clr = forma.fill._xPr.find(qn("a:solidFill"))[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(valor * 1000)))


def alfa_linea(forma, valor):
    ln = forma.line._get_or_add_ln()
    clr = ln.find(qn("a:solidFill"))[0]
    a = etree.SubElement(clr, qn("a:alpha"))
    a.set("val", str(int(valor * 1000)))


def sombra(forma, blur=28, dist=10, opacidad=45):
    spPr = forma._element.spPr
    efectos = etree.SubElement(spPr, qn("a:effectLst"))
    s = etree.SubElement(efectos, qn("a:outerShdw"), blurRad=str(Pt(blur)), dist=str(Pt(dist)), dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(s, qn("a:srgbClr"), val="06170D")
    etree.SubElement(c, qn("a:alpha"), val=str(opacidad * 1000))


def caja(s, x, y, w, h, color=BLANCO, transp=12, borde=MENTA, borde_transp=55, radio=0.08, con_sombra=True):
    """Fondo de color para un bloque. Los de vidrio blanco ya no se dibujan: el texto va
    directo sobre el fondo, como en la portada."""
    if color == BLANCO:
        return None
    f = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    f.adjustments[0] = radio
    f.fill.solid()
    f.fill.fore_color.rgb = rgb(color)
    if transp < 100:
        alfa(f, transp)
    if borde:
        f.line.color.rgb = rgb(borde)
        f.line.width = Pt(1.25)
        alfa_linea(f, borde_transp)
    else:
        f.line.fill.background()
    if con_sombra:
        sombra(f)
    f.text_frame.text = ""
    return f


def texto(s, x, y, w, h, contenido, tam=16, color=BLANCO, fuente=CUERPO, negrita=False, alinear=PP_ALIGN.LEFT,
          anclar=MSO_ANCHOR.TOP, interlineado=1.15, espaciado=0, mayus=False):
    """contenido: str o lista de párrafos; cada párrafo puede ser str o lista de (texto, {opciones})."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = anclar
    parrafos = contenido if isinstance(contenido, list) else [contenido]
    for i, p in enumerate(parrafos):
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = alinear
        par.line_spacing = interlineado
        trozos = p if isinstance(p, list) else [(p, {})]
        for t, op in trozos:
            r = par.add_run()
            r.text = t.upper() if (op.get("mayus", mayus)) else t
            fnt = r.font
            fnt.name = op.get("fuente", fuente)
            fnt.size = Pt(op.get("tam", tam))
            fnt.bold = op.get("negrita", negrita)
            fnt.italic = op.get("cursiva", False)
            fnt.color.rgb = rgb(op.get("color", color))
            if espaciado:
                r._r.get_or_add_rPr().set("spc", str(espaciado * 100))
        if isinstance(p, list) and p and p[0][1].get("espacio_despues"):
            par.space_after = Pt(p[0][1]["espacio_despues"])
    return tb


def imagen(s, nombre, x, y, w=None, h=None):
    return s.shapes.add_picture(rec(nombre), Inches(x), Inches(y),
                                Inches(w) if w else None, Inches(h) if h else None)


def encabezado(s, etiqueta, titulo, ancho=10.2, tam_titulo=38):
    texto(s, 0.85, 0.55, ancho, 0.35, etiqueta, tam=13, color=LIMA, negrita=True, espaciado=3, mayus=True)
    # Josefin Sans es ancha: se achica el título para que quepa en una línea.
    tam_titulo = min(tam_titulo, int(ancho * 72 / (len(titulo) * 0.6)))
    texto(s, 0.85, 0.9, ancho, 1.1, titulo, tam=tam_titulo, fuente=TITULO, negrita=True, color=BLANCO, interlineado=1.0)
    linea = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.88), Inches(1.95), Inches(0.9), Inches(0.07))
    linea.fill.solid()
    linea.fill.fore_color.rgb = rgb(LIMA)
    linea.line.fill.background()


def pie(s, fuente=None):
    t = f"No estamos solos LMH  ·  {numero[0]:02d}"
    if fuente:
        t = f"{fuente}  ·  " + t
    claro = fondo_actual[0] == "fondo_brillante.jpg"
    texto(s, 0.85, 6.95, 10.5, 0.3, t, tam=10, color=OSCURO if claro else TEXTO_SUAVE)
    imagen(s, "escudo_liceo.png", 12.35, 6.82, h=0.48)


def deco(s, piezas):
    """Elementos 3D decorativos: lista de (archivo, x, y, tamaño)."""
    for nombre, x, y, tam in piezas:
        imagen(s, nombre, x, y, w=tam, h=tam)


def notas(s, t):
    s.notes_slide.notes_text_frame.text = t


def lineas_estimadas(t, tam, ancho, factor=0.55):
    return max(1, math.ceil(len(t) * factor * tam / 72 / max(ancho, 0.5)))


def tarjeta(s, x, y, w, h, titulo, cuerpo, numero_tarjeta=None, acento=False, tam_titulo=19, tam_cuerpo=14):
    """Bloque sin cuadro: un objeto 3D como viñeta, el título y el texto flotando sobre el fondo."""
    claro = fondo_actual[0] == "fondo_brillante.jpg"
    tarjeta.n = getattr(tarjeta, "n", 0) + 1
    if numero_tarjeta is not None:
        texto(s, x - 0.03, y - 0.1, 1.4, 0.75, numero_tarjeta, tam=34, fuente=TITULO, negrita=True, color=OLIVA if claro else MENTA)
    else:
        marca = "obj_estrella.png" if acento else ["esfera_lima.png", "esfera_menta.png", "esfera_blanca.png"][tarjeta.n % 3]
        imagen(s, marca, x - 0.06, y - 0.08, w=0.62, h=0.62)
    yt = y + 0.6
    alto_t = lineas_estimadas(titulo, tam_titulo, w - 0.1, 0.6) * tam_titulo * 1.08 / 72
    texto(s, x, yt, w - 0.1, alto_t + 0.1, titulo, tam=tam_titulo, fuente=TITULO, negrita=True,
          color=OSCURO if claro else LIMA, interlineado=1.0)
    texto(s, x, yt + alto_t + 0.12, w - 0.15, max(0.5, h - 0.8 - alto_t), cuerpo, tam=tam_cuerpo,
          color=OSCURO if claro else TEXTO_SUAVE, interlineado=1.22)


def barra(s, x, y, h, color):
    """Línea vertical de color para separar un bloque sin encerrarlo."""
    b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    b.adjustments[0] = 0.5
    b.fill.solid()
    b.fill.fore_color.rgb = rgb(color)
    b.line.fill.background()


def borde_celda(celda, lado, color=None, alfa_valor=40, ancho=1):
    tcPr = celda._tc.get_or_add_tcPr()
    ln = etree.SubElement(tcPr, qn(f"a:ln{lado}"), w=str(Pt(ancho)))
    if color:
        sf = etree.SubElement(ln, qn("a:solidFill"))
        c = etree.SubElement(sf, qn("a:srgbClr"), val=color)
        etree.SubElement(c, qn("a:alpha"), val=str(alfa_valor * 1000))
    else:
        etree.SubElement(ln, qn("a:noFill"))


def tabla(s, x, y, w, filas, anchos, tam=12.5, alto_fila=0.42, resaltar=None):
    """Tabla sin relleno: encabezado lima, texto claro y una línea fina menta entre filas."""
    nf, nc = len(filas), len(filas[0])
    forma = s.shapes.add_table(nf, nc, Inches(x), Inches(y), Inches(w), Inches(alto_fila * nf))
    gt = forma.table
    tblPr = forma._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    estilo = tblPr.find(qn("a:tableStyleId"))
    if estilo is None:
        estilo = etree.SubElement(tblPr, qn("a:tableStyleId"))
    estilo.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # Sin estilo, sin cuadrícula
    for j, a in enumerate(anchos):
        gt.columns[j].width = Inches(a)
    for i, fila in enumerate(filas):
        gt.rows[i].height = Inches(alto_fila)
        for j, valor in enumerate(fila):
            celda = gt.cell(i, j)
            celda.fill.background()
            celda.margin_left = Inches(0.02 if j == 0 else 0.12)
            celda.margin_right = Inches(0.08)
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = celda.text_frame
            tf.text = ""
            r = tf.paragraphs[0].add_run()
            r.text = valor
            destacada = resaltar is not None and i == resaltar
            r.font.name = TITULO if i == 0 else CUERPO
            r.font.size = Pt(tam + (1 if i == 0 else 0))
            r.font.bold = i == 0 or destacada or j == 0
            r.font.color.rgb = rgb(LIMA if (i == 0 or destacada) else (BLANCO if j == 0 else TEXTO_SUAVE))
            for lado in ("L", "R", "T"):
                borde_celda(celda, lado)
            borde_celda(celda, "B", LIMA if i == 0 else MENTA, 90 if i == 0 else 30, 1.5 if i == 0 else 0.75)
    return gt


def flecha(s, x, y, w=0.45, h=0.3, color=LIMA):
    f = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    f.fill.solid()
    f.fill.fore_color.rgb = rgb(color)
    f.line.fill.background()
    return f


def pastilla(s, x, y, w, t, color=LIMA, color_texto=OSCURO, tam=15, h=0.5):
    f = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    f.adjustments[0] = 0.5
    f.fill.solid()
    f.fill.fore_color.rgb = rgb(color)
    f.line.fill.background()
    sombra(f, blur=16, dist=5, opacidad=35)
    tf = f.text_frame
    tf.margin_left = tf.margin_right = Inches(0.12)
    tf.text = ""
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = t
    r.font.name, r.font.size, r.font.bold = CUERPO, Pt(tam), True
    r.font.color.rgb = rgb(color_texto)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    return f


# ---------------------------------------------------------------------------
# 1. Portada
# ---------------------------------------------------------------------------
s = nueva("fondo_liceo.jpg")
deco(s, [("esfera_lima.png", 5.6, 0.25, 1.3), ("esfera_menta.png", 6.7, 5.7, 0.9), ("anillo_menta.png", 0.1, 5.4, 1.6)])
imagen(s, "escudo_liceo.png", 0.9, 0.6, h=1.15)
texto(s, 2.2, 0.75, 6, 0.9, ["Liceo Militar de Honduras", "Proyecto tecnológico 2026"], tam=13, color=LIMA, negrita=True, espaciado=2, mayus=True)
texto(s, 0.85, 2.05, 7.2, 2.4, ["No estamos", "solos LMH"], tam=66, fuente=TITULO, negrita=True, color=BLANCO, interlineado=0.95)
texto(s, 0.9, 4.45, 6.6, 1.0, "Plataforma web anónima de apoyo emocional para los estudiantes del liceo", tam=20, color=TEXTO_SUAVE, interlineado=1.2)
pastilla(s, 0.9, 5.65, 2.6, "Escucha · Apoyo · Esperanza", tam=13)
texto(s, 0.9, 6.45, 9, 0.5, "Presenta: Padilla Garay   ·   Asesor(a): [nombre del asesor]   ·   [fecha de la exposición]", tam=12.5, color=TEXTO_SUAVE)
imagen(s, "tigrito_psicologo.png", 8.35, 1.25, h=6.25)
notas(s, "Saludo al jurado. Presentar el proyecto \"No estamos solos LMH\": una plataforma web donde los estudiantes del Liceo Militar de Honduras pueden pedir ayuda emocional sin dar su nombre. Presentar a Tigrito, la mascota del programa.")

# ---------------------------------------------------------------------------
# 2. El problema
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="corazon")
encabezado(s, "El problema", "Muchos estudiantes no piden ayuda a tiempo")
for i, (t, c) in enumerate([
    ("Miedo a ser identificados", "Temen que se enteren sus compañeros o su familia, y prefieren callar."),
    ("Una psicóloga, horario limitado", "Fuera del horario no hay a quién acudir, ni forma de saber qué caso es urgente."),
    ("Señales que llegan tarde", "Sin un canal, una señal de crisis puede no llegar nunca a quien puede ayudar."),
]):
    tarjeta(s, 0.85 + i * 3.95, 2.35, 3.65, 2.55, t, c, numero_tarjeta=f"0{i + 1}")
caja(s, 0.85, 5.25, 11.55, 1.25, color=LIMA, transp=100, borde=None)
texto(s, 1.2, 5.38, 2.3, 1.0, "15 %", tam=48, fuente=TITULO, negrita=True, color=OSCURO, anclar=MSO_ANCHOR.MIDDLE)
texto(s, 3.3, 5.42, 8.9, 1.0, "de los adolescentes de 10 a 19 años de América Latina y el Caribe vive con un trastorno mental diagnosticado, según UNICEF.",
      tam=16, color=OSCURO, anclar=MSO_ANCHOR.MIDDLE)
pie(s, "Fuente del dato: UNICEF")
notas(s, "Explicar el problema con ejemplos del liceo: el miedo a ser identificado, una sola psicóloga con horario limitado y señales de crisis que pueden llegar tarde. El dato de UNICEF da contexto regional.")

# ---------------------------------------------------------------------------
# 3. Necesidad (diapositiva de una idea)
# ---------------------------------------------------------------------------
s = nueva("fondo_brillante.jpg")
deco(s, [("esfera_blanca.png", 8.3, 0.45, 1.25), ("capsula.png", 7.55, 5.75, 1.45), ("esfera_lima.png", 0.4, 6.05, 0.8)])
imagen(s, "obj_burbuja.png", 11.35, 1.0, w=1.55, h=1.55)
imagen(s, "tigrito_psicologo.png", 8.55, 1.75, w=5.0, h=5.0)
texto(s, 0.95, 0.85, 7.5, 0.4, "Necesidad tecnológica identificada", tam=14, color=OSCURO, negrita=True, espaciado=3, mayus=True)
texto(s, 0.95, 1.45, 7.6, 3.4, "Un canal anónimo, disponible las 24 horas, que detecte señales de riesgo y avise a la psicóloga a tiempo.",
      tam=34, fuente=TITULO, negrita=True, color=OSCURO, interlineado=1.05)
for i, t in enumerate(["Anónimo, sin nombre", "Siempre disponible", "Alertas al instante"]):
    pastilla(s, 0.95 + i * 2.5, 5.0, 2.35, t, color=OSCURO, color_texto=LIMA, tam=13, h=0.6)
pie(s)
notas(s, "La necesidad en una frase. Los tres requisitos guían todo el diseño: anonimato, disponibilidad y alertas en tiempo real.")

# ---------------------------------------------------------------------------
# 4. Referentes
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="estrella")
encabezado(s, "Referentes", "Nos basamos en Yana y Psonríe")
tabla(s, 0.85, 2.35, 11.6, [
    ["Aspecto", "Yana", "Psonríe", "No estamos solos LMH"],
    ["Qué es", "App de acompañamiento emocional (México, 2017)", "Psicología online con psicólogos colegiados (España, 2016)", "Plataforma web del liceo"],
    ["Quién atiende", "Chatbot 24/7", "Psicólogo por chat, voz o video", "Psicóloga del liceo y asistente Sami, Leo o Nia"],
    ["Costo", "Gratis, con suscripción opcional", "Pago por sesión", "Gratis para el estudiante"],
    ["Lo que tomamos", "Chat 24/7, diario de ánimo, ejercicios", "Atención sin cita, privacidad", "Todo eso, enfocado en el liceo y con alertas de riesgo"],
], [2.0, 3.1, 3.1, 3.4], tam=12.5, alto_fila=0.78)
pie(s, "Fuentes: prensa sobre Yana y ficha de Psonríe")
notas(s, "Yana es una app mexicana con chatbot, registro de ánimo y ejercicios. Psonríe es un servicio español de psicología en línea. Tomamos lo mejor de ambas y lo adaptamos al liceo: gratis, anónimo y conectado con la psicóloga real.")

# ---------------------------------------------------------------------------
# 5. Propuesta
# ---------------------------------------------------------------------------
s = nueva("fondo_3.jpg", icono="burbuja")
encabezado(s, "La propuesta", "Una plataforma para hablar, aprender y cuidarse")
deco(s, [("anillo_lima.png", 11.3, 5.75, 1.3)])
for i, (t, c) in enumerate([
    ("Hablar", "Chat anónimo con la psicóloga del liceo y con el asistente Sami, Leo o Nia, que hace preguntas de seguimiento."),
    ("Aprender", "64 artículos, un glosario de 26 palabras de psicología y 24 herramientas con sus pasos."),
    ("Cuidarse", "Diario de ánimo, 3 autoevaluaciones, un mural interactivo y un muro para compartir."),
]):
    imagen(s, ["esfera_lima.png", "esfera_menta.png", "esfera_blanca.png"][i], 1.0 + i * 3.95, 2.35, w=1.2, h=1.2)
    texto(s, 1.15 + i * 3.95, 3.55, 3.1, 0.7, t, tam=30, fuente=TITULO, negrita=True, color=LIMA)
    texto(s, 1.15 + i * 3.95, 4.3, 3.15, 1.9, c, tam=15, color=TEXTO_SUAVE, interlineado=1.25)
pie(s)
notas(s, "La propuesta en tres verbos: hablar, aprender y cuidarse.")

# ---------------------------------------------------------------------------
# 6. Funcionalidades
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="engranaje")
encabezado(s, "Funcionalidades", "Lo que hace la plataforma")
items = [
    ("Chat con la psicóloga", "Anónimo y en tiempo real."),
    ("Asistente Sami, Leo o Nia", "Pregunta, recuerda el tema y sugiere herramientas."),
    ("Detección de riesgo", "8 categorías; alerta a la psicóloga."),
    ("Líneas de ayuda", "911 y Teléfono de la Esperanza 150."),
    ("Herramientas", "24 ejercicios con pasos, como la respiración 4-4-4-4."),
    ("Infórmate", "64 artículos, 4 secciones y glosario."),
    ("Diario y tests", "Ánimo diario, estrés y hábitos de estudio."),
    ("Panel de la psicóloga", "Conversaciones, alertas y estadísticas."),
]
for k, (t, c) in enumerate(items):
    fila, col = divmod(k, 4)
    tarjeta(s, 0.85 + col * 2.95, 2.35 + fila * 2.2, 2.75, 2.0, t, c, acento=(k == 7), tam_titulo=16, tam_cuerpo=13)
pie(s)
notas(s, "Recorrer las ocho funciones. La detección de riesgo revisa cada mensaje con 8 categorías; si encuentra señales, el mensaje no lo responde el asistente: se deriva a la psicóloga y se muestran las líneas 150 y 911.")

# ---------------------------------------------------------------------------
# 7. Conoce a Tigrito
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg")
deco(s, [("esfera_lima.png", 6.0, 0.4, 1.0), ("esfera_menta.png", 12.0, 5.9, 0.9)])
foto = imagen(s, "tigrito_uniforme.jpg", 6.75, 1.0, w=5.9)
sombra(foto, blur=30, dist=12, opacidad=50)
texto(s, 0.85, 0.85, 5.5, 0.4, "Identidad del programa", tam=13, color=LIMA, negrita=True, espaciado=3, mayus=True)
texto(s, 0.85, 1.3, 5.6, 1.6, ["Conoce a", "Tigrito"], tam=50, fuente=TITULO, negrita=True, color=BLANCO, interlineado=0.95)
texto(s, 0.85, 3.05, 5.4, 1.6, "La mascota del programa: un tigre del liceo, en uniforme y como psicólogo, que acompaña a los estudiantes en toda la página.",
      tam=16, color=TEXTO_SUAVE, interlineado=1.25)
imagen(s, "tigres_bravios.png", 0.9, 4.65, h=1.95)
imagen(s, "escudo_liceo.png", 2.7, 4.7, h=1.85)
imagen(s, "informatica.png", 4.55, 4.6, h=2.0)
texto(s, 6.75, 5.1, 5.9, 0.5, "Tigrito frente al Liceo Militar de Honduras", tam=12, color=TEXTO_SUAVE, alinear=PP_ALIGN.CENTER)
pie(s)
notas(s, "Tigrito es la identidad del programa: aparece en la portada de la página y en las herramientas. Usa los símbolos del liceo: el escudo, los Tigres Bravíos y el área de Informática.")

# ---------------------------------------------------------------------------
# 8. Diseño e interactividad
# ---------------------------------------------------------------------------
s = nueva("fondo_3.jpg", icono="estrella")
encabezado(s, "Diseño e interactividad", "Una página que invita a quedarse")
items = [
    ("Mural de psicología", "Respira conmigo (4-4-4-4), mito o realidad y datos de psicología."),
    ("Libros de apoyo", "Un estante con 7 libros reales y para qué ayuda cada uno."),
    ("Recursos ilustrados", "15 temas de psicología, cada uno con su dibujo."),
    ("Contadores reales", "Visitas y conversaciones, contadas de forma anónima."),
    ("Muro interactivo", "Experiencias anónimas; las crisis no se publican."),
    ("Pensado para todos", "Celular y computadora; las animaciones se detienen si piden menos movimiento."),
]
for k, (t, c) in enumerate(items):
    fila, col = divmod(k, 3)
    tarjeta(s, 0.85 + col * 3.95, 2.35 + fila * 2.2, 3.65, 2.0, t, c, acento=(k == 5), tam_titulo=18, tam_cuerpo=13.5)
pie(s, "Paleta: #1F4B2C · #4D7111 · #C3E956 · #91EAAF · #F4FFFC")
notas(s, "La experiencia: mural con actividades, libros de apoyo, recursos ilustrados, contadores reales y un muro anónimo. Todo funciona en celular y computadora con la paleta verde del proyecto.")

# ---------------------------------------------------------------------------
# 9. Crece con el liceo (novedades)
# ---------------------------------------------------------------------------
s = nueva("fondo_brillante.jpg", icono="rayo")
deco(s, [("anillo_menta.png", 0.2, 5.6, 1.5)])
texto(s, 0.95, 0.75, 10, 0.4, "Novedades", tam=14, color=OSCURO, negrita=True, espaciado=3, mayus=True)
texto(s, 0.95, 1.2, 10.5, 1.1, "Una página que crece con el liceo", tam=36, fuente=TITULO, negrita=True, color=OSCURO)
for i, (t, c) in enumerate([
    ("La psicóloga la actualiza", "Agrega o esconde artículos, herramientas, frases e historias, y modera el muro. Sin tocar código."),
    ("Cuentas opcionales", "Apodo y avatar; el diario saluda al estudiante y lo sigue en cualquier celular. Los chats siguen anónimos."),
    ("App instalable", "Se instala en el celular con el ícono del liceo y muestra las líneas 150 y 911 sin internet."),
]):
    tarjeta(s, 0.95 + i * 3.95, 2.85, 3.6, 3.5, t, c, tam_titulo=22, tam_cuerpo=15.5)
pie(s)
notas(s, "Lo más nuevo: la psicóloga puede cambiar el contenido desde su panel, sin programar; los estudiantes pueden crear una cuenta opcional con apodo y avatar; y la página se instala como aplicación en el celular.")

# ---------------------------------------------------------------------------
# 10. Esquema del sistema
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="nube")
encabezado(s, "Esquema del sistema", "Quién la usa y con qué se conecta")
cols = [("Usuarios", ["Estudiante: anónimo, cuenta opcional", "Psicóloga: usuario y contraseña"], False),
        ("Plataforma web", ["Páginas, chat en tiempo real, filtro de riesgo, asistente y panel de contenido"], True),
        ("Servicios", ["Base de datos SQLite", "API de Claude (IA)", "YouTube y Spotify"], False)]
for i, (t, lineas, ac) in enumerate(cols):
    x = 0.85 + i * 4.15
    if ac:
        caja(s, x, 2.6, 3.55, 3.2, color=LIMA, transp=100, borde=None)
    else:
        caja(s, x, 2.6, 3.55, 3.2)
    texto(s, x + 0.3, 2.85, 3.0, 0.6, t, tam=22, fuente=TITULO, negrita=True, color=OSCURO if ac else LIMA)
    texto(s, x + 0.3, 3.55, 3.0, 2.2, lineas, tam=15, color=OSCURO if ac else TEXTO_SUAVE, interlineado=1.35)
    if i < 2:
        flecha(s, x + 3.66, 4.05, 0.42, 0.32)
pie(s)
notas(s, "A la izquierda los usuarios; en el centro la plataforma web; a la derecha los servicios con los que se conecta.")

# ---------------------------------------------------------------------------
# 11. Arquitectura
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="engranaje")
encabezado(s, "Diagrama de arquitectura", "Cliente, servidor y datos")
bloques = [("Cliente (navegador)", ["HTML5 y CSS3", "JavaScript", "Cliente Socket.IO", "App instalable (PWA)"], False, 3.3),
           ("Servidor (Python + Flask)", ["Rutas y eventos en tiempo real", "Filtro de riesgo", "Asistente y respaldo por reglas", "Contenido, cuentas y panel"], True, 3.9),
           ("Datos y servicios", ["SQLite", "API de Claude (IA)", "YouTube y Spotify"], False, 3.0)]
x = 0.85
for i, (t, lineas, ac, w) in enumerate(bloques):
    barra(s, x + 0.05, 2.65, 3.1, LIMA if ac else MENTA)
    texto(s, x + 0.25, 2.7, w - 0.4, 0.6, t, tam=18, fuente=TITULO, negrita=True, color=LIMA)
    texto(s, x + 0.25, 3.35, w - 0.4, 2.4, ["•  " + l for l in lineas], tam=14, color=TEXTO_SUAVE, interlineado=1.45)
    x += w
    if i < 2:
        flecha(s, x + 0.12, 4.05, 0.36, 0.28)
        texto(s, x - 0.05, 4.4, 0.7, 0.6, ["HTTPS", "WebSocket"] if i == 0 else ["Consultas", "y API"], tam=9, color=LIMA, alinear=PP_ALIGN.CENTER)
        x += 0.62
pie(s, "Arquitectura cliente-servidor con tiempo real")
notas(s, "Arquitectura cliente-servidor. El navegador se comunica por HTTPS y WebSocket con el servidor en Python y Flask, que guarda los datos en SQLite y consulta la IA por la API de Claude.")

# ---------------------------------------------------------------------------
# 12. Base de datos
# ---------------------------------------------------------------------------
s = nueva("fondo_3.jpg", icono="candado")
encabezado(s, "Base de datos", "Tablas en SQLite, sin nombres de estudiantes")
tabla(s, 0.85, 2.3, 11.6, [
    ["Tabla", "Para qué sirve", "Campos principales"],
    ["sesiones y mensajes", "Cada conversación anónima y sus mensajes", "token, emisor, contenido, nivel de riesgo"],
    ["animos", "Diario de ánimo", "token, fecha, estado"],
    ["usuarios", "Cuentas opcionales (apodo y avatar)", "apodo, contraseña cifrada, avatar"],
    ["resultados_test", "Nivel general de cada autoevaluación", "test, nivel, fecha"],
    ["muro_posts y resenas", "Muro y opiniones (se pueden esconder)", "contenido, fecha, oculto"],
    ["contenido_*", "Lo que edita la psicóloga", "artículos, herramientas, textos, ajustes"],
    ["contactos_emergencia", "Número que se deja en una crisis (opcional)", "token, teléfono, plataforma"],
], [3.0, 4.6, 4.0], tam=12.5, alto_fila=0.52)
pie(s, "Cada estudiante es solo un código anónimo (token)")
notas(s, "Ninguna tabla guarda nombres ni correos. Cada conversación se identifica con un código anónimo. Las cuentas usan un apodo y la contraseña se guarda cifrada. La base de datos nunca se sube a GitHub.")

# ---------------------------------------------------------------------------
# 13. Flujo de datos
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="burbuja")
encabezado(s, "Flujo de datos", "Qué pasa con un mensaje del estudiante")
pasos = ["El estudiante escribe", "Se guarda con su código anónimo", "El filtro mide el riesgo", "Llega la respuesta en tiempo real"]
for i, p in enumerate(pasos):
    x = 0.85 + i * 3.0
    caja(s, x, 2.4, 2.55, 1.75)
    texto(s, x + 0.25, 2.5, 1, 0.6, str(i + 1), tam=30, fuente=TITULO, negrita=True, color=LIMA)
    texto(s, x + 0.25, 3.1, 2.15, 1.0, p, tam=15, color=BLANCO, interlineado=1.15)
    if i < 3:
        flecha(s, x + 2.6, 3.12, 0.34, 0.26)
barra(s, 0.9, 4.6, 1.7, MENTA)
texto(s, 1.15, 4.65, 5.1, 0.5, "Riesgo bajo", tam=20, fuente=TITULO, negrita=True, color=MENTA)
texto(s, 1.15, 5.2, 5.1, 1.2, "Responde el asistente: pregunta, recuerda el tema y sugiere herramientas.", tam=14.5, color=BLANCO, interlineado=1.25)
barra(s, 6.85, 4.6, 1.7, "F08A7A")
texto(s, 7.1, 4.65, 5.1, 0.5, "Riesgo medio, alto o crítico", tam=20, fuente=TITULO, negrita=True, color="FFC2B8")
texto(s, 7.1, 5.2, 5.1, 1.2, "Mensaje de apoyo, alerta a la psicóloga y, si es crítico, botones del 150 y el 911.", tam=14.5, color=BLANCO, interlineado=1.25)
pie(s, "Una crisis nunca la responde el asistente")
notas(s, "El recorrido de un mensaje y la bifurcación clave: si hay riesgo, el mensaje no va al asistente; la psicóloga recibe una alerta y el estudiante ve las líneas 150 y 911.")

# ---------------------------------------------------------------------------
# 14. Procesos
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="escudo")
encabezado(s, "Diagrama de procesos", "Del primer mensaje a la atención")
filas = [("Estudiante", ["Entra sin registrarse", "Elige: psicóloga, asistente o herramientas", "Escribe o usa una herramienta"], "vidrio"),
         ("Sistema", ["¿El mensaje tiene señales de riesgo? Si las tiene, avisa a la psicóloga al instante."], "lima"),
         ("Psicóloga", ["Recibe la alerta o el mensaje", "Responde desde su panel", "Da seguimiento y actualiza el contenido"], "oscuro")]
for f, (actor, etapas, estilo) in enumerate(filas):
    y = 2.4 + f * 1.45
    texto(s, 0.85, y + 0.3, 1.9, 0.5, actor, tam=17, fuente=TITULO, negrita=True, color=LIMA)
    if len(etapas) == 1:
        caja(s, 2.8, y, 9.6, 1.1, color=LIMA, transp=100, borde=None)
        texto(s, 3.1, y + 0.1, 9.1, 0.9, etapas[0], tam=15.5, negrita=True, color=OSCURO, anclar=MSO_ANCHOR.MIDDLE)
        continue
    for k, e in enumerate(etapas):
        x = 2.8 + k * 3.3
        barra(s, x, y + 0.15, 0.8, MENTA if estilo == "oscuro" else LIMA)
        texto(s, x + 0.2, y + 0.08, 2.5, 0.95, e, tam=14, color=BLANCO, anclar=MSO_ANCHOR.MIDDLE)
        if k < 2:
            flecha(s, x + 2.9, y + 0.42, 0.34, 0.26)
pie(s)
notas(s, "Tres actores: el estudiante, el sistema y la psicóloga. El sistema revisa cada mensaje y avisa al instante si hay riesgo.")

# ---------------------------------------------------------------------------
# 15. Stack
# ---------------------------------------------------------------------------
s = nueva("fondo_3.jpg", icono="engranaje")
encabezado(s, "Stack tecnológico", "Lenguajes, frameworks y herramientas")
tabla(s, 0.85, 2.3, 11.6, [
    ["Capa", "Tecnología", "Para qué"],
    ["Lenguajes", "Python 3.11, HTML5, CSS3, JavaScript", "Lógica y páginas"],
    ["Framework web", "Flask 3.0", "Rutas y plantillas"],
    ["Tiempo real", "Flask-SocketIO 5.3 (WebSocket)", "Chat y alertas al instante"],
    ["Base de datos", "SQLite", "Mensajes anónimos, cuentas y contenido"],
    ["Inteligencia artificial", "API de Claude (librería anthropic)", "Asistente que conversa"],
    ["Animaciones y app", "Canvas, WebGL, SVG y PWA", "Mural, efectos y app instalable"],
    ["Producción y control", "Gunicorn, Eventlet, Git, GitHub, Render", "Publicar y guardar versiones"],
], [3.0, 4.8, 3.8], tam=12.5, alto_fila=0.52)
pie(s, "Todo el software es libre o gratuito, salvo la API de IA")
notas(s, "Python y Flask en el servidor, Socket.IO para el tiempo real, SQLite como base de datos, la API de Claude para la IA, y Canvas, WebGL y SVG para las animaciones. La página es una PWA: se instala como app.")

# ---------------------------------------------------------------------------
# 16. Tiempo real
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="rayo")
encabezado(s, "Red y tiempo real", "Lo que necesita para funcionar en vivo")
for i, (t, lineas) in enumerate([
    ("Cómo funciona", ["Una sala de Socket.IO por conversación", "Una sala para el panel de la psicóloga", "Eventos: mensaje, alerta de riesgo, estado de la psicóloga y apoyo de emergencia"]),
    ("Qué necesita", ["Servidor encendido las 24 horas (hosting)", "HTTPS con certificado de seguridad", "Internet en el liceo y en el celular de la psicóloga", "Clave de la IA y contraseñas en el servidor, nunca en el código"]),
]):
    x = 0.85 + i * 5.9
    barra(s, x, 2.6, 3.6, LIMA if i == 0 else MENTA)
    texto(s, x + 0.35, 2.65, 5, 0.6, t, tam=22, fuente=TITULO, negrita=True, color=LIMA)
    texto(s, x + 0.35, 3.35, 5, 3.0, ["•  " + l for l in lineas], tam=15, color=TEXTO_SUAVE, interlineado=1.4)
pie(s)
notas(s, "Cada conversación es una sala privada; la psicóloga tiene una sala propia para las alertas. Para funcionar necesita un servidor encendido, HTTPS, internet y las claves configuradas en el servidor.")

# ---------------------------------------------------------------------------
# 17. Demostración
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="estrella")
encabezado(s, "Demostración funcional", "Muestra en vivo: 7 minutos")
demo = [("Portada: logo, carrusel, mural de psicología y línea 150", "1 min"),
        ("Herramientas: la respiración 4-4-4-4, con sus pasos a la vista", "1 min"),
        ("Chat con Sami: pregunta, recuerda el tema y sugiere una herramienta", "2 min"),
        ("Mensaje de prueba con riesgo: alerta en el panel de la psicóloga", "2 min"),
        ("Panel de contenido, cuenta opcional e instalar la app", "1 min")]
for i, (t, m) in enumerate(demo):
    y = 2.35 + i * 0.86
    caja(s, 0.85, y, 11.6, 0.72, transp=10)
    texto(s, 1.1, y + 0.06, 0.7, 0.6, str(i + 1), tam=24, fuente=TITULO, negrita=True, color=LIMA, anclar=MSO_ANCHOR.MIDDLE)
    texto(s, 1.8, y + 0.06, 9.2, 0.6, t, tam=16, color=BLANCO, anclar=MSO_ANCHOR.MIDDLE)
    pastilla(s, 11.2, y + 0.14, 1.05, m, tam=12, h=0.44)
pie(s, "La prueba de riesgo se hace en un entorno de prueba")
notas(s, "Guion de la demostración. En el paso 3, escribir \"estoy estresado por los exámenes\" y luego \"desde la semana pasada\": el asistente recuerda el tema y sugiere la respiración 4-4-4-4. En el paso 4, tener abiertas dos ventanas: el panel de la psicóloga y el chat del estudiante.")

# ---------------------------------------------------------------------------
# 18. Los cuatro modelos
# ---------------------------------------------------------------------------
s = nueva("fondo_3.jpg", icono="reloj")
encabezado(s, "Los cuatro modelos", "Tecnología, gestión, cronograma y cierre")
for k, (t, c) in enumerate([
    ("Tecnología", "Aplicación web cliente-servidor, con chat en tiempo real, asistente con respaldo por reglas y app instalable."),
    ("Gestión", "Desarrollo iterativo e incremental: cada mejora es un cambio guardado en Git (76 hasta hoy)."),
    ("Cronograma", "Fases cortas; al final de cada una se prueba en el navegador antes de seguir."),
    ("Cierre", "Pruebas finales, publicación en internet, entrega al liceo y capacitación de la psicóloga."),
]):
    fila, col = divmod(k, 2)
    tarjeta(s, 0.85 + col * 5.9, 2.35 + fila * 2.15, 5.6, 1.95, t, c, numero_tarjeta=f"0{k + 1}", tam_titulo=20, tam_cuerpo=14)
pie(s)
notas(s, "Los cuatro modelos que guían el proyecto: tecnológico, de gestión, de cronograma y de cierre.")

# ---------------------------------------------------------------------------
# 19. Presupuesto
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="moneda")
encabezado(s, "Presupuesto", "Hardware, licencias y servicios")
tabla(s, 0.85, 2.35, 11.6, [
    ["Concepto", "Detalle", "Costo"],
    ["Software", "Python, Flask, SQLite y Git: licencias libres", "L 0"],
    ["GitHub", "Repositorio del código", "L 0"],
    ["Computadora", "Equipo propio para desarrollar", "L 0"],
    ["Celular de la psicóloga", "Para recibir avisos y WhatsApp (ref. Galaxy A17)", "L 4,939"],
    ["Internet", "Red del liceo", "L 0"],
    ["Hosting, dominio e IA", "Pago mensual: ver la cotización", "siguiente"],
], [3.4, 6.4, 1.8], tam=13.5, alto_fila=0.56)
pie(s, "Precio de referencia: tienda Samsung Honduras; verificar antes de comprar")
notas(s, "Todo el software es libre. El único gasto de hardware recomendado es un celular para la psicóloga, para que reciba los avisos sin usar su teléfono personal.")

# ---------------------------------------------------------------------------
# 20. Hosting
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="nube")
encabezado(s, "Cotización de hosting y servicios", "Dónde publicar la plataforma")
tabla(s, 0.85, 2.3, 11.6, [
    ["Servicio", "Precio", "Chat en vivo", "Nota"],
    ["Render Free", "US$0", "Sí", "Para mostrarla: se duerme y no guarda datos"],
    ["Render Starter", "US$7/mes ≈ L 188", "Sí", "Recomendado para uso real"],
    ["Railway Hobby", "US$5/mes + uso ≈ L 134", "Sí", "El precio sube con el uso"],
    ["PythonAnywhere", "US$10/mes ≈ L 268", "No", "No admite WebSockets"],
    ["Dominio .com", "US$11 el primer año ≈ L 295", "No aplica", "Renovación: US$15 a 19 al año"],
    ["IA: Claude Opus 5.5", "≈ US$36/mes ≈ L 965", "No aplica", "Unos 3,000 mensajes al mes"],
    ["IA: Claude Haiku 5.5", "≈ US$1/mes ≈ L 27", "No aplica", "Opción más económica"],
], [2.7, 3.0, 1.7, 4.2], tam=12, alto_fila=0.5, resaltar=2)
pie(s, "Precios de referencia 2026, L 26.80 por dólar (BCH)")
notas(s, "Para la exposición basta el plan gratis de Render. Para uso real recomendamos Render Starter: siempre encendido, admite WebSockets y guarda los datos con su disco.")

# ---------------------------------------------------------------------------
# 21. Inversión
# ---------------------------------------------------------------------------
s = nueva("fondo_brillante.jpg", icono="moneda")
deco(s, [("esfera_lima.png", 0.35, 6.1, 0.8)])
texto(s, 0.95, 0.75, 10, 0.4, "Inversión del proyecto", tam=14, color=OSCURO, negrita=True, espaciado=3, mayus=True)
texto(s, 0.95, 1.2, 10.5, 1.0, "Costo estimado del primer año", tam=36, fuente=TITULO, negrita=True, color=OSCURO)
for i, (t, cifra, det) in enumerate([
    ("Con IA Claude Opus 5.5", "L 19,063", "Celular L 4,939 + hosting, dominio e IA US$527 (≈ L 14,124)"),
    ("Con IA Claude Haiku 5.5", "L 7,807", "Celular L 4,939 + hosting, dominio e IA US$107 (≈ L 2,868)"),
]):
    x = 0.95 + i * 5.85
    barra(s, x, 2.75, 3.3, OSCURO if i == 0 else OLIVA)
    texto(s, x + 0.4, 2.85, 4.8, 0.5, t, tam=17, negrita=True, color=OLIVA)
    texto(s, x + 0.4, 3.35, 5.0, 1.3, cifra, tam=72, fuente=TITULO, negrita=True, color=OSCURO)
    texto(s, x + 0.4, 4.85, 4.8, 1.2, det, tam=15, color=OSCURO, interlineado=1.25)
pie(s, "Estimado con 3,000 mensajes al mes y L 26.80 por dólar")
notas(s, "La inversión del primer año depende del modelo de IA. Desde el segundo año ya no se compra el celular. Sin IA, el asistente sigue funcionando con respuestas preparadas.")

# ---------------------------------------------------------------------------
# 22. Gantt
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="reloj")
encabezado(s, "Diagrama de Gantt", "Cronograma de ejecución, 2026")
columnas = ["Antes", "2–3 oct", "4–5 oct", "7–8 oct", "12 oct", "13–15 oct", "Próximo"]
fases = ["Análisis y planificación", "Base: chat, asistente y riesgo", "Funciones: diario, tests, Infórmate",
         "Diseño e identidad del liceo", "IA, herramientas y línea 150", "Interactividad, paleta y contenidos", "Pruebas, publicación y entrega"]
x0, y0, ancho_col, alto = 4.15, 2.3, 1.18, 0.5
for j, c in enumerate(columnas):
    texto(s, x0 + j * ancho_col, y0, ancho_col, 0.4, c, tam=11.5, negrita=True, color=LIMA, alinear=PP_ALIGN.CENTER)
for i, fase in enumerate(fases):
    y = y0 + 0.5 + i * 0.6
    texto(s, 0.85, y + 0.05, 3.25, 0.45, fase, tam=13, color=BLANCO, anclar=MSO_ANCHOR.MIDDLE)
    for j in range(len(columnas)):
        activa = (i == j)
        b = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x0 + j * ancho_col + 0.05), Inches(y + 0.04), Inches(ancho_col - 0.1), Inches(alto - 0.06))
        b.adjustments[0] = 0.3
        b.fill.solid()
        if activa and i < 6:
            b.fill.fore_color.rgb = rgb(LIMA)
            b.line.fill.background()
            sombra(b, blur=10, dist=3, opacidad=35)
        elif activa:
            b.fill.fore_color.rgb = rgb(MENTA)
            alfa(b, 30)
            b.line.color.rgb = rgb(MENTA)
            b.line.dash_style = 4
        else:
            b.fill.fore_color.rgb = rgb(BLANCO)
            alfa(b, 8)
            b.line.fill.background()
pie(s, "Lima: fase hecha (fechas reales del historial de Git) · Menta: pendiente · [fechas de planificación por confirmar]")
notas(s, "El Gantt usa las fechas reales de Git. Agregar las fechas exactas de la planificación. Falta la última fase: pruebas, publicación y entrega al liceo.")

# ---------------------------------------------------------------------------
# 23. Cumplimiento de fases
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="escudo")
encabezado(s, "Cumplimiento de fases", "6 de 7 fases cumplidas")
tabla(s, 0.85, 2.3, 11.6, [
    ["Fase", "Estado", "Evidencia"],
    ["Análisis y planificación", "Cumplida", "Problema y necesidad definidos"],
    ["Desarrollo base", "Cumplida", "Primera versión, 2 de octubre"],
    ["Funciones", "Cumplida", "Diario, tests, Infórmate y herramientas"],
    ["Diseño institucional", "Cumplida", "Colores del liceo, escudo y uso en celular"],
    ["IA y contenidos", "Cumplida", "Asistente y línea 150, 12 de octubre"],
    ["Interactividad y paleta", "Cumplida", "Mural, recursos ilustrados, panel y app, 13 a 15 de octubre"],
    ["Pruebas y publicación", "En curso", "Falta publicar en Render y activar la clave de IA"],
], [3.6, 1.9, 6.1], tam=13, alto_fila=0.52, resaltar=7)
pie(s, "76 cambios guardados en Git")
notas(s, "Seis de siete fases terminadas, con evidencia en Git. Falta publicar en Render, activar la clave de la IA y hacer las pruebas finales con la psicóloga.")

# ---------------------------------------------------------------------------
# 24. Impacto
# ---------------------------------------------------------------------------
s = nueva("fondo_3.jpg", icono="corazon")
encabezado(s, "Impacto y aporte tecnológico", "Lo que gana el Liceo Militar de Honduras")
for k, (t, c) in enumerate([
    ("Un canal anónimo, 24/7", "Los estudiantes piden ayuda sin miedo a ser identificados."),
    ("Alertas para la psicóloga", "Los casos con riesgo le llegan primero, al instante."),
    ("Datos sin nombres", "Estadísticas para decidir, sin exponer a nadie."),
    ("Bajo costo y fácil de mantener", "Software libre; la psicóloga actualiza el contenido sin programar."),
]):
    fila, col = divmod(k, 2)
    tarjeta(s, 0.85 + col * 5.9, 2.35 + fila * 2.15, 5.6, 1.95, t, c, tam_titulo=20, tam_cuerpo=15)
pie(s)
notas(s, "El aporte a la institución: canal anónimo, alertas a tiempo, estadísticas sin nombres y bajo costo.")

# ---------------------------------------------------------------------------
# 25. Conclusiones
# ---------------------------------------------------------------------------
s = nueva("fondo_1.jpg", icono="estrella")
encabezado(s, "Conclusiones y logros", "Una plataforma real, lista para publicarse")
for i, (cifra, t) in enumerate([("64", "artículos en Infórmate"), ("24", "herramientas con pasos"), ("8", "categorías de riesgo"), ("26", "conceptos en el glosario")]):
    x = 0.85 + i * 2.95
    caja(s, x, 2.35, 2.75, 2.1)
    texto(s, x + 0.3, 2.5, 2.3, 1.1, cifra, tam=58, fuente=TITULO, negrita=True, color=LIMA)
    texto(s, x + 0.3, 3.65, 2.3, 0.7, t, tam=14.5, color=TEXTO_SUAVE)
caja(s, 0.85, 4.8, 11.6, 1.55, color=LIMA, transp=100, borde=None)
texto(s, 1.2, 4.9, 11, 1.35, "La tecnología acerca la ayuda a los estudiantes sin reemplazar a la psicóloga: la apoya y le avisa a tiempo cuando alguien la necesita.",
      tam=20, fuente=TITULO, negrita=True, color=OSCURO, anclar=MSO_ANCHOR.MIDDLE, interlineado=1.1)
pie(s)
notas(s, "Logros en números. La conclusión principal: la tecnología no reemplaza a la psicóloga, la apoya.")

# ---------------------------------------------------------------------------
# 26. Recomendaciones y aprendizajes
# ---------------------------------------------------------------------------
s = nueva("fondo_2.jpg", icono="burbuja")
encabezado(s, "Recomendaciones y aprendizajes", "Lo que sigue y lo que aprendimos")
barra(s, 0.9, 2.6, 3.7, LIMA)
texto(s, 1.2, 2.55, 5, 0.6, "Recomendaciones", tam=22, fuente=TITULO, negrita=True, color=LIMA)
texto(s, 1.2, 3.2, 5.1, 3.2, ["•  " + l for l in ["Que la psicóloga revise las frases de riesgo", "Publicar con HTTPS y contraseñas propias",
                                                    "Respaldar la base de datos cada semana", "Probarla con un grupo de estudiantes",
                                                    "Definir un protocolo de crisis del liceo", "Usar historias reales, anónimas y con permiso"]],
      tam=14.5, color=TEXTO_SUAVE, interlineado=1.4)
barra(s, 6.9, 2.6, 3.7, MENTA)
texto(s, 7.2, 2.55, 5, 0.6, "Aprendizajes técnicos", tam=22, fuente=TITULO, negrita=True, color=LIMA)
texto(s, 7.2, 3.2, 5.1, 3.2, ["•  " + l for l in ["Chat en tiempo real con WebSockets", "Privacidad desde el diseño", "IA con límites y respaldo",
                                                    "Diseño accesible para celular", "Contenido editable sin programar", "Control de versiones con Git"]],
      tam=14.5, color=BLANCO, interlineado=1.4)
pie(s)
notas(s, "Recomendaciones para el liceo y aprendizajes técnicos del proyecto.")

# ---------------------------------------------------------------------------
# 27. Gracias
# ---------------------------------------------------------------------------
s = nueva("fondo_liceo_der.jpg")
deco(s, [("esfera_lima.png", 0.6, 0.5, 1.1), ("anillo_menta.png", 0.3, 5.5, 1.6), ("esfera_menta.png", 5.4, 6.0, 0.8)])
imagen(s, "tigrito_psicologo.png", 0.9, 1.1, h=6.2)
caja(s, 6.55, 1.15, 6.25, 5.0, color=OSCURO, transp=82, borde=MENTA, borde_transp=55)
texto(s, 6.9, 1.4, 6, 1.6, "Gracias", tam=80, fuente=TITULO, negrita=True, color=LIMA)
texto(s, 6.95, 2.95, 5.8, 0.7, "¿Preguntas?", tam=30, fuente=TITULO, color=BLANCO)
texto(s, 6.95, 3.9, 5.9, 1.2, ["noestamossoloslmh@gmail.com", "Teléfono 3174-6810 · San Pedro Sula, Honduras"], tam=15, color=TEXTO_SUAVE, interlineado=1.35)
pastilla(s, 6.95, 5.15, 5.4, "Si necesitas hablar con alguien: llama gratis al 150", tam=14.5, h=0.62)
imagen(s, "escudo_liceo.png", 11.7, 6.25, h=0.95)
notas(s, "Agradecer al jurado y abrir el espacio de preguntas. Cerrar con el mensaje del proyecto: nadie tiene que pasar por esto solo. Mencionar la línea gratuita 150.")

prs.save(SALIDA)
print("Diapositivas:", len(prs.slides), "->", SALIDA)
